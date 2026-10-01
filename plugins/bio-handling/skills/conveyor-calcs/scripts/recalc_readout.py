#!/usr/bin/env python3
"""Recalculate a JMS calc workbook in LibreOffice (on a throwaway copy) and read results.

Usage:
  python3 recalc_readout.py <filled workbook> [--cells M19 P40 AU58 ...] [--sheet Design]
                            [--outputs] [--json out.json]

  --cells    list of cells to print (value after recalculation)
  --outputs  also list every cell styled as an output (pass / warning / failure)
             with its label, per the Instructions color legend

The original file is never modified, and the LibreOffice-saved copy must NOT be
delivered: LibreOffice rewrites the package (macros, images and formatting can
change). It is only for reading numbers. Excel recalculates the delivered file on
open (fill_calc.py sets fullCalcOnLoad).

Known LibreOffice gaps in the JMS templates: XLOOKUP-based part-number cells
return #NAME? in LibreOffice; Excel computes them. Treat #NAME? in part-number
cells as expected, not as a calc failure.
"""
import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from inspect_template import read_legend, color_key, label_for  # noqa: E402

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
Sub RecalcSave()
  ThisComponent.calculateAll()
  ThisComponent.calculateAll()
  ThisComponent.store()
  ThisComponent.close(True)
End Sub
</script:module>"""


def recalc_copy(src):
    tmp = Path(tempfile.mkdtemp(prefix="jmsrecalc_"))
    dst = tmp / ("calc" + Path(src).suffix)
    shutil.copy(src, dst)
    prof = tmp / "profile"
    url = prof.as_uri()
    env = dict(os.environ, HOME=str(tmp))
    subprocess.run(["soffice", "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
                   capture_output=True, timeout=120, env=env)
    mdir = prof / "user" / "basic" / "Standard"
    mdir.mkdir(parents=True, exist_ok=True)
    (mdir / "Module1.xba").write_text(MACRO)
    subprocess.run(["soffice", "--headless", f"-env:UserInstallation={url}",
                    "macro:///Standard.Module1.RecalcSave", str(dst)],
                   capture_output=True, timeout=300, env=env)
    return dst


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--cells", nargs="*", default=[])
    ap.add_argument("--sheet", default="Design")
    ap.add_argument("--outputs", action="store_true")
    ap.add_argument("--json")
    a = ap.parse_args()

    calc = recalc_copy(a.workbook)
    wbv = openpyxl.load_workbook(calc, data_only=True)
    ws = wbv[a.sheet]
    res = {}
    for ref in a.cells:
        res[ref] = ws[ref].value
        print(f"{ref:>6} = {ws[ref].value}")
    if a.outputs:
        wbf = openpyxl.load_workbook(a.workbook)
        legend = read_legend(wbf)
        inv = {v: k for k, v in legend.items()}
        wsf = wbf[a.sheet]
        merged = {}
        for mr in wsf.merged_cells.ranges:
            for r in range(mr.min_row, mr.max_row + 1):
                for c in range(mr.min_col, mr.max_col + 1):
                    merged[(r, c)] = (mr.min_row, mr.min_col)
        print("\nOutput cells (class, label, value after recalc):")
        for row in wsf.iter_rows():
            for c in row:
                cls = inv.get(color_key(c))
                if cls not in ("output_pass", "output_warning", "output_failure"):
                    continue
                if (c.row, c.column) in merged and merged[(c.row, c.column)] != (c.row, c.column):
                    continue
                v = ws[c.coordinate].value
                res[c.coordinate] = v
                print(f"{c.coordinate:>6} {cls:<15} {label_for(wsf, c.row, c.column, merged)[:45]:<45} = {v}")
    if a.json:
        json.dump(res, open(a.json, "w"), indent=1, default=str)
    shutil.rmtree(calc.parent, ignore_errors=True)


if __name__ == "__main__":
    main()
