#!/usr/bin/env python3
"""Write the project selection log (<Project Number>-Selection Log.xlsx) from a JSON file.

  python3 selection_log.py log.json "<out dir>/25026-Selection Log.xlsx"

log.json:
{
  "project": "25026", "prepared_by": "ASB", "submittal_rev": "A", "date": "2026-10-05",
  "components": [ {"item": 1, "equipment": "Zero Speed Switch", "source": "spec",
                   "reference": "46 21 73 2.6.H", "tags": "SSC-1, SSC-2", "qty": 2,
                   "vendor": "Electro-Sensors", "model": "SCP1000 115 VAC 800-020100",
                   "library_file": "Zero Speed Switch/Electro Sensors/...Cat Cut.pdf",
                   "pages": "2", "output_pdf": "25026-Zero Speed Switch.pdf",
                   "status": "final"} ],
  "compliance": [ {"item": 1, "requirement": "115 VAC supply", "reference": "2.6.H.2",
                   "offered": "SCP1000 115 VAC (p2)", "result": "compliant", "note": ""} ],
  "decisions":  [ {"item": 1, "deviation": "...", "options": "...", "decision": "exception",
                   "decided_by": "ASB", "date": "2026-10-06", "reference": "Exception 3"} ],
  "revisions":  [ {"rev": "A", "date": "2026-10-06", "change": "Initial submittal"} ]
}
source: spec | calc | jms-std.  status: final | provisional | missing data sheet.
result: compliant | deviation | n/a.  decision: accept | RFI | exception | open.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

SHEETS = {
    "Components": [("item", "Item", 6), ("equipment", "Equipment", 24), ("source", "Source", 10),
                   ("reference", "Spec para / calc cell", 22), ("tags", "Tags", 18), ("qty", "Qty", 6),
                   ("vendor", "Vendor", 18), ("model", "Model / part number", 30),
                   ("library_file", "Library file", 50), ("pages", "Marked pages", 12),
                   ("output_pdf", "Output PDF", 32), ("status", "Status", 16)],
    "Compliance": [("item", "Item", 6), ("requirement", "Requirement", 40), ("reference", "Spec para", 16),
                   ("offered", "Cut sheet offers (file, page)", 40), ("result", "Result", 12), ("note", "Note", 40)],
    "Decisions": [("item", "Item", 6), ("deviation", "Deviation", 40), ("options", "Options considered", 40),
                  ("decision", "Decision", 12), ("decided_by", "Decided by", 12), ("date", "Date", 12),
                  ("reference", "RFI / exception ref", 20)],
    "Revisions": [("rev", "Submittal rev", 12), ("date", "Date", 12), ("change", "What changed", 60)],
}
FLAG = {"deviation": "FCE4D6", "provisional": "FFF2CC", "missing data sheet": "FCE4D6", "open": "FFF2CC"}


def main(src, out):
    data = json.load(open(src))
    wb = Workbook()
    wb.remove(wb.active)
    for name, cols in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append([f"Project {data.get('project', '')} - Cut sheet selection log - {name}"])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([f"Prepared by {data.get('prepared_by', '')}   Submittal rev {data.get('submittal_rev', '')}   "
                   f"{data.get('date', '')}"])
        ws.append([c[1] for c in cols])
        for c in ws[3]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        rows = data.get(name.lower(), [])
        for r in rows:
            ws.append([r.get(k, "") for k, _, _ in cols])
            fill = next((FLAG[str(v).lower()] for v in r.values() if str(v).lower() in FLAG), None)
            if fill:
                for c in ws[ws.max_row]:
                    c.fill = PatternFill("solid", fgColor=fill)
        for i, (_, _, w) in enumerate(cols, start=1):
            ws.column_dimensions[ws.cell(row=3, column=i).column_letter].width = w
        for row in ws.iter_rows(min_row=4):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A4"
        if rows:
            ref = f"A3:{ws.cell(row=3, column=len(cols)).column_letter}{ws.max_row}"
            t = Table(displayName=name, ref=ref)
            t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
            ws.add_table(t)
    wb.save(out)
    counts = {n: len(data.get(n.lower(), [])) for n in SHEETS}
    print(f"Wrote {out}: " + ", ".join(f"{k} {v}" for k, v in counts.items()))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
